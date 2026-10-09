"""CSV / Excel export helpers."""
from __future__ import annotations

import csv
from pathlib import Path


def export_csv(path: str | Path, headers: list[str], rows: list[list]) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)
    return p


def export_excel(path: str | Path, headers: list[str], rows: list[list], title: str = "Report") -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    # autosize columns, freeze header, add filter
    widths = [len(str(h)) for h in headers]
    for r in rows:
        for i, v in enumerate(r):
            if i < len(widths):
                widths[i] = max(widths[i], len(str(v)))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 12), 50)
    ws.freeze_panes = "A2"
    try:
        ws.auto_filter.ref = ws.dimensions
    except Exception:
        pass
    wb.save(p)
    return p
